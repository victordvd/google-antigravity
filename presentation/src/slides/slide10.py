# slide10.py — Settings Deep Dive: Models Quota and Usage Management

from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
import theme as T
import primitives as P

def build(prs, slide):
    P.add_background(slide)
    P.add_eyebrow(slide, 2, "INTERFACE")

    title_segments = [
        ("後台設定：模型配額與算力掌控", {"color": T.ORANGE}),
        ("：掌握 Gemini 週期額度與費用邊界", {}),
    ]
    P.add_argument_title(slide, title_segments, y=T.TITLE_Y)

    support_segments = [
        ("依任務複雜度切換模型，透過後台即時掌握", {}),
        ("5 小時與每週使用額度", {"color": T.INK, "bold": True}),
        ("，並開啟超額防護避免額外費用。", {}),
    ]
    P.add_supporting_sentence(slide, support_segments, y=Inches(1.58))
    P.add_hairline(slide, Inches(2.2))

    card_y = Inches(2.45)
    card_h = Inches(4.55)
    img_w = Inches(7.4)
    col_gap = Inches(0.28)
    info_w = Inches(4.41)

    # ------------------------------------------------------------- Left: Large Annotated Screenshot
    img_x = T.MARGIN_X
    img_path = T.ASSETS_DIR / "ui-img" / "settings-models.png"
    # settings-models.png is 1172 x 807 -> scale = 0.005638 -> placed w = 6.61", h = 4.55"
    P.add_screenshot(slide, str(img_path), img_x, card_y, max_w=img_w, max_h=card_h, frame=True)

    # --- Callout 2: Enable AI Credit Overages Toggle Highlight Box (y ~ 295..385 in 807)
    box2_x = img_x + Inches(1.97)
    box2_y = card_y + Inches(1.66)
    box2_w = Inches(4.31)
    box2_h = Inches(0.51)
    box2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, box2_x, box2_y, box2_w, box2_h)
    box2.fill.background()
    box2.line.color.rgb = T.RED
    box2.line.width = Pt(2.0)
    P.no_shadow(box2)
    P.add_status_pill(slide, box2_x + Inches(0.06), box2_y - Inches(0.28), "② 費用超額防護開關 (預設關閉)",
                      bg=T.RED, fg=T.ON_DARK, w=Inches(2.55), h=Inches(0.24), size=10)

    # --- Callout 1: Gemini Models Quota Highlight Box (y ~ 420..600 in 807)
    box1_x = img_x + Inches(1.97)
    box1_y = card_y + Inches(2.37)
    box1_w = Inches(4.31)
    box1_h = Inches(1.02)
    box1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, box1_x, box1_y, box1_w, box1_h)
    box1.fill.background()
    box1.line.color.rgb = T.ORANGE
    box1.line.width = Pt(2.25)
    P.no_shadow(box1)
    P.add_status_pill(slide, box1_x + box1_w - Inches(2.35), box1_y + Inches(0.06), "① Gemini 雙週期額度監控",
                      bg=T.ORANGE, fg=T.ON_DARK, w=Inches(2.25), h=Inches(0.24), size=10)

    # ------------------------------------------------------------- Right: Companion Explanation Panel
    info_x = img_x + img_w + col_gap
    P.add_panel(slide, info_x, card_y, info_w, card_h, fill=T.PANEL_BG, line=T.PANEL_LN)

    # Header
    tb_h, tf_h = P.textbox(slide, info_x + Inches(0.2), card_y + Inches(0.18), info_w - Inches(0.4), Inches(0.28))
    P.rich_par(tf_h, [
        ("MODELS & USAGE  ·  ", {"color": T.ORANGE, "mono": True, "bold": True, "size": 11}),
        ("算力與額度策略", {"color": T.MUTED, "size": 11}),
    ], first=True)

    tb_t, tf_t = P.textbox(slide, info_x + Inches(0.2), card_y + Inches(0.48), info_w - Inches(0.4), Inches(0.38))
    P.rich_par(tf_t, [("模型選擇與成本邊界把關", {"color": T.INK, "bold": True, "size": 16})], first=True)

    P.add_hairline(slide, card_y + Inches(0.92), x=info_x + Inches(0.2), w=info_w - Inches(0.4))

    # Notes inside a single contiguous TextFrame to prevent overlapping
    sections = [
        ("① 雙重週期額度掌控", [
            ("5 小時重置額度", "應對連續密集對話與多步驟測試，每 5 小時自動刷新。"),
            ("每週總量額度", "宏觀掌握長期算力使用進度，避免單一專案耗盡額度。"),
        ]),
        ("② Flash vs. Pro 彈性分工", [
            ("Gemini 3.8 Flash", "速度極快、耗額低，推薦用於磁碟整理、文字摘要。"),
            ("Gemini 3.8 Pro", "深層規劃與複雜邏輯除錯，適合跨檔案大綱與爬蟲分析。"),
        ]),
        ("③ 帳單安全隔離", [
            ("關閉 Credit Overages", "用量達上限時自動暫停，杜絕非預期之雲端信用卡扣款。"),
        ]),
    ]

    tb_notes, tf_notes = P.textbox(slide, info_x + Inches(0.2), card_y + Inches(1.04),
                                   info_w - Inches(0.4), card_h - Inches(1.15))
    first_par = True
    for s_idx, (s_title, s_points) in enumerate(sections):
        p_title = P.rich_par(tf_notes, [(s_title, {"color": T.ORANGE, "bold": True, "size": 12.5})],
                             first=first_par, space_after=4)
        if s_idx > 0:
            p_title.space_before = Pt(8)
        first_par = False

        for j, (pname, pdesc) in enumerate(s_points):
            P.rich_par(tf_notes, [
                (f"▸ {pname}：", {"color": T.INK, "bold": True, "size": 10.5}),
                (pdesc, {"color": T.INK_SOFT, "size": 10.5}),
            ], space_after=3 if j < len(s_points) - 1 else 0, line=1.18)
