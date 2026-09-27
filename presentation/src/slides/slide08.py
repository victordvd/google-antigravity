# slide08.py — UI Deep Dive: Full Antigravity 2.0 Desktop View with Callouts

from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
import theme as T
import primitives as P

def build(prs, slide):
    P.add_background(slide)
    P.add_eyebrow(slide, 2, "INTERFACE")

    title_segments = [
        ("原生介面全貌導覽", {"color": T.ORANGE}),
        ("：導航側欄、對話畫布與提示詞輸入框", {}),
    ]
    P.add_argument_title(slide, title_segments, y=T.TITLE_Y)

    support_segments = [
        ("真實桌面環境一覽：", {}),
        ("左側專案管理、中央發布目標", {"color": T.INK, "bold": True}),
        ("，底部輸入列直接調用快捷指令與切換模型。", {}),
    ]
    P.add_supporting_sentence(slide, support_segments, y=Inches(1.58))
    P.add_hairline(slide, Inches(2.2))

    card_y = Inches(2.45)
    card_h = Inches(4.55)
    img_w = Inches(8.5)
    col_gap = Inches(0.28)
    info_w = Inches(3.31)

    # ------------------------------------------------------------- Left: Large Annotated Screenshot
    img_x = T.MARGIN_X
    img_path = T.ASSETS_DIR / "ui-img" / "antigravity.png"
    P.add_screenshot(slide, str(img_path), img_x, card_y, max_w=img_w, max_h=card_h, frame=True)

    # --- Callout 1: Sidebar Highlight Box
    sb_x = img_x + Inches(0.06)
    sb_y = card_y + Inches(0.33)
    sb_w = Inches(1.50)
    sb_h = Inches(4.00)
    box1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, sb_x, sb_y, sb_w, sb_h)
    box1.fill.background()
    box1.line.color.rgb = T.ORANGE
    box1.line.width = Pt(2.0)
    P.no_shadow(box1)
    P.add_status_pill(slide, sb_x + Inches(0.06), sb_y + Inches(0.06), "① 左側邊欄",
                      bg=T.ORANGE, fg=T.ON_DARK, w=Inches(1.05), h=Inches(0.24), size=10)

    # --- Callout 2: Input Bar Highlight Box
    in_x = img_x + Inches(2.84)
    in_y = card_y + Inches(2.06)
    in_w = Inches(4.25)
    in_h = Inches(0.60)
    box2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, in_x, in_y, in_w, in_h)
    box2.fill.background()
    box2.line.color.rgb = T.ORANGE
    box2.line.width = Pt(2.0)
    P.no_shadow(box2)
    P.add_status_pill(slide, in_x + Inches(0.06), in_y - Inches(0.28), "② 提示詞輸入列 (/ 與 @ 快捷列)",
                      bg=T.ORANGE, fg=T.ON_DARK, w=Inches(2.65), h=Inches(0.24), size=10)

    # ------------------------------------------------------------- Right: Companion Explanation Panel
    info_x = img_x + img_w + col_gap
    P.add_panel(slide, info_x, card_y, info_w, card_h, fill=T.PANEL_BG, line=T.PANEL_LN)

    # Header
    tb_h, tf_h = P.textbox(slide, info_x + Inches(0.2), card_y + Inches(0.18), info_w - Inches(0.4), Inches(0.28))
    P.rich_par(tf_h, [
        ("UI BREAKDOWN  ·  ", {"color": T.ORANGE, "mono": True, "bold": True, "size": 11}),
        ("視窗區域解析", {"color": T.MUTED, "size": 11}),
    ], first=True)

    tb_t, tf_t = P.textbox(slide, info_x + Inches(0.2), card_y + Inches(0.48), info_w - Inches(0.4), Inches(0.38))
    P.rich_par(tf_t, [("核心控制項位置對照", {"color": T.INK, "bold": True, "size": 16})], first=True)

    P.add_hairline(slide, card_y + Inches(0.92), x=info_x + Inches(0.2), w=info_w - Inches(0.4))

    # Notes inside a single contiguous TextFrame to prevent overlapping
    sections = [
        ("① 導航側欄 (Sidebar)", [
            ("New Conversation", "開闢乾淨獨立脈絡，杜絕歷史任務污染。"),
            ("Projects (agy)", "快速切換本機工作空間與研究資料夾。"),
            ("Scheduled Tasks", "設定定時背景自動化排程任務。"),
        ]),
        ("② 提示詞輸入列 (Prompt Bar)", [
            ("白話指派", "自然語言描述目標與格式需求。"),
            ("模型切換", "即時切換 Gemini Flash / Pro。"),
            ("斜線與引用", "鍵入 / 叫出計畫，鍵入 @ 注入本機資料。"),
        ]),
    ]

    tb_notes, tf_notes = P.textbox(slide, info_x + Inches(0.2), card_y + Inches(1.04),
                                   info_w - Inches(0.4), card_h - Inches(1.15))
    first_par = True
    for s_idx, (s_title, s_points) in enumerate(sections):
        p_title = P.rich_par(tf_notes, [(s_title, {"color": T.ORANGE, "bold": True, "size": 12.5})],
                             first=first_par, space_after=4)
        if s_idx > 0:
            p_title.space_before = Pt(10)
        first_par = False

        for j, (pname, pdesc) in enumerate(s_points):
            P.rich_par(tf_notes, [
                (f"▸ {pname}：", {"color": T.INK, "bold": True, "size": 10.5}),
                (pdesc, {"color": T.INK_SOFT, "size": 10.5}),
            ], space_after=3 if j < len(s_points) - 1 else 0, line=1.18)
